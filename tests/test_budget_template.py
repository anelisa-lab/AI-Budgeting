"""
The downloadable budget spreadsheet (app/budget_template.py).

These read the generated workbook back with openpyxl and check its structure and
formulas. The numbers the formulas produce were checked separately by
recalculating a filled-in copy in LibreOffice (see docs/BUDGET_CYCLES_AND_TEMPLATE.md).
"""

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.budget_template import (
    LOG_FIRST_ROW,
    build_budget_workbook,
    template_filename,
    tracker_columns,
)

D = Decimal
START = date(2026, 9, 29)
CATS = [("Groceries", D("600")), ("Toiletries", None), ("Transport", D("250"))]


def build(days=30, cats=CATS, total="1715", pct="10", start=START):
    data = build_budget_workbook(cats, D(total), days, start, D(pct))
    return load_workbook(io.BytesIO(data))


def header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=30):
        if row[0].value == "Category":
            return row[0].row
    raise AssertionError("no header row")


# --- week or day columns ---------------------------------------------------

def test_a_month_gets_weekly_columns_ending_on_payout_day():
    kind, cols = tracker_columns(START, 30)
    assert kind == "week"
    assert [c.label for c in cols] == ["Week 1", "Week 2", "Week 3", "Week 4", "Week 5"]
    assert cols[0].start == START
    assert cols[0].end == START + timedelta(days=6)
    assert cols[-1].end == START + timedelta(days=30)         # payout day
    for a, b in zip(cols, cols[1:]):
        assert b.start == a.end + timedelta(days=1)            # no gaps, no overlap


def test_two_weeks_gets_two_weekly_columns():
    kind, cols = tracker_columns(START, 14)
    assert kind == "week" and len(cols) == 2
    assert cols[-1].end == START + timedelta(days=14)


def test_a_week_gets_a_column_per_day_including_payout_day():
    kind, cols = tracker_columns(START, 7)
    assert kind == "day"
    assert len(cols) == 8                                       # start day .. payout day
    assert cols[0].label == "Tue 29 Sep"
    assert cols[-1].label == "Payout day"
    assert all(c.start == c.end for c in cols)


def test_a_zero_day_period_is_refused():
    with pytest.raises(ValueError):
        tracker_columns(START, 0)


# --- the Budget sheet ------------------------------------------------------

def test_workbook_has_a_budget_sheet_and_a_spend_log():
    wb = build()
    assert wb.sheetnames == ["Budget", "Spend log"]


def test_the_money_block_uses_formulas_not_typed_results():
    ws = build()["Budget"]
    assert ws["B4"].value == 1715.0
    assert ws["B5"].value == pytest.approx(0.10)
    assert ws["B6"].value == "=ROUND(B4*B5,2)"
    assert ws["B7"].value == "=B4-B6"
    assert ws["B10"].value == "=B8+B9"
    assert ws["B11"].value == "=B9+1"          # start day and payout day both count
    assert ws["B12"].value == "=IF(B11>0,B7/B11,0)"
    assert ws["B8"].value.date() == START
    assert ws["B9"].value == 30


def test_the_day_count_matches_the_apps_daily_split():
    """Template and dashboard must agree on "R a day" (payout day counts)."""
    from app.budget_split import days_remaining

    wb = build(days=30)
    ws = wb["Budget"]
    # B11 is "=B9+1"; the split's count from the start day to the payout date is the same.
    assert days_remaining(START, START + timedelta(days=30)) == 31 == ws["B9"].value + 1


def test_chosen_categories_are_the_rows_in_order():
    ws = build()["Budget"]
    head = header_row(ws)
    assert [ws.cell(row=head + i, column=1).value for i in (1, 2, 3)] == ["Groceries", "Toiletries", "Transport"]
    assert ws.cell(row=head + 4, column=1).value.startswith("Other")
    assert ws.cell(row=head + 5, column=1).value == "Total"


def test_planned_amounts_are_prefilled_only_when_given():
    ws = build()["Budget"]
    head = header_row(ws)
    assert ws.cell(row=head + 1, column=2).value == 600.0
    assert ws.cell(row=head + 2, column=2).value is None
    assert ws.cell(row=head + 3, column=2).value == 250.0


def test_a_month_has_a_per_week_column_and_a_week_does_not():
    month = build(days=30)["Budget"]
    week = build(days=7)["Budget"]
    head_m, head_w = header_row(month), header_row(week)
    assert "Planned per week (R)" in [c.value for c in month[head_m]]
    assert "Planned per week (R)" not in [c.value for c in week[head_w]]


def test_spent_columns_total_the_log_by_category_and_date_range():
    ws = build(days=30)["Budget"]
    head = header_row(ws)
    # Column F is Week 1 in the monthly layout (A..E fixed, F.. weeks).
    formula = ws.cell(row=head + 1, column=6).value
    assert formula.startswith("=SUMIFS('Spend log'!$D$")
    assert "$A%d" % (head + 1) in formula                       # matches its own category cell
    assert '">="&F$14' in formula and '"<="&F$15' in formula    # the week's from/to dates
    assert f"$D${LOG_FIRST_ROW}:$D$1000" in formula             # the example row above is not counted


def test_the_last_week_ends_on_the_payout_date_cell():
    ws = build(days=30)["Budget"]
    assert ws["J15"].value == "=$B$10"
    assert ws["F14"].value == "=$B$8"
    assert ws["G14"].value == "=F15+1"


def test_left_and_status_columns_handle_an_unplanned_category():
    ws = build(days=30)["Budget"]
    head = header_row(ws)
    left = ws.cell(row=head + 2, column=12).value              # Toiletries, "Left"
    assert left.startswith('=IF(B') and '"",""' in left         # blank until a plan is typed
    status = ws.cell(row=head + 2, column=13).value
    assert "Plan me" in status and "Over budget" in status


def test_the_total_row_sums_every_category_and_other():
    ws = build(days=30)["Budget"]
    head = header_row(ws)
    total = head + 5
    assert ws.cell(row=total, column=2).value == f"=SUM(B{head + 1}:B{head + 3})"
    assert ws.cell(row=total, column=6).value == f"=SUM(F{head + 1}:F{head + 4})"   # includes Other


def test_not_planned_yet_is_available_minus_planned():
    ws = build()["Budget"]
    head = header_row(ws)
    gap = head + 7
    assert ws.cell(row=gap, column=1).value == "Not planned yet (R)"
    assert ws.cell(row=gap, column=2).value == f"=B7-B{head + 5}"


def test_a_legend_tells_the_student_what_to_fill_in():
    ws = build()["Budget"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "How to use this" in text
    assert "yellow" in text.lower()
    assert "Spend log" in text


# --- the Spend log ---------------------------------------------------------

def test_spend_log_has_headers_an_uncounted_example_and_a_category_dropdown():
    wb = build()
    log = wb["Spend log"]
    assert [log.cell(row=3, column=c).value for c in (1, 2, 3, 4)] == [
        "Date", "What I bought", "Category", "Amount (R)"]
    assert "example" in log["B4"].value.lower()
    assert log["C4"].value == "Groceries"
    assert LOG_FIRST_ROW == 5                                   # counted rows start below the example
    validations = log.data_validations.dataValidation
    lists = [v for v in validations if v.type == "list"]
    assert len(lists) == 1
    assert lists[0].formula1 == "=Budget!$A$17:$A$19"           # exactly the three chosen categories
    assert "C5:C1000" in str(lists[0].sqref)


# --- inputs ----------------------------------------------------------------

def test_duplicate_and_messy_names_are_cleaned():
    ws = build(cats=[(" groceries ", None), ("Groceries", D("5")), ("=Evil", None)])["Budget"]
    head = header_row(ws)
    assert ws.cell(row=head + 1, column=1).value == "groceries"
    assert ws.cell(row=head + 2, column=1).value == "Evil"      # no formula injection
    assert not str(ws.cell(row=head + 2, column=1).value).startswith("=")


def test_no_categories_is_refused():
    with pytest.raises(ValueError):
        build_budget_workbook([], D("1000"), 30, START, D("0"))


def test_filenames_say_which_kind_of_budget():
    assert template_filename(START, 7) == "uniwallet-weekly-budget-2026-09-29.xlsx"
    assert template_filename(START, 14) == "uniwallet-fortnightly-budget-2026-09-29.xlsx"
    assert template_filename(START, 30) == "uniwallet-monthly-budget-2026-09-29.xlsx"

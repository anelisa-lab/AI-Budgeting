"""
Budget spreadsheet template — the "plan my priorities" download.

The student picks the categories that matter this cycle; this module builds an
.xlsx they can fill in on their phone or laptop. No database, no FastAPI: give
it the money, the period and the categories, get bytes back, so every number in
it is unit-tested against a real workbook (tests/test_budget_template.py).

WHAT THE WORKBOOK CONTAINS
--------------------------
"Budget"     the money at the top (received, savings, what is spendable, days),
             then one row per chosen category: a yellow "Planned" cell to fill
             in, its share, its per-day (and per-week) amount, and how much has
             been SPENT in each week — or each day, for a budget that lasts a
             week or less — pulled from the log, and what is left.
"Spend log"  Date / What / Category / Amount. The student types spends here and
             the Budget sheet totals them by category and by week or day.

WEEK OR MONTH
-------------
A budget that lasts up to a week gets one column per DAY; anything longer gets
one column per WEEK. Either way the columns end on payout day, and the daily
figure divides by the same day count as the app's Daily Budget Split
(app/budget_split.days_remaining: start day AND payout day both count), so the
template and the dashboard never disagree about "R a day".

Everything that can be a formula is one, so the sheet recalculates when the
student changes the money or a planned amount.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal
from typing import List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.budget_categories import dedupe_categories

FONT = "Arial"
LOG_LAST_ROW = 1000          # rows of the Spend log the Budget sheet reads
LOG_FIRST_ROW = 5            # row 4 is the (uncounted) example row

MONEY = '"R"#,##0.00;[Red]-"R"#,##0.00;"–"'
PERCENT = '0.0%;[Red]-0.0%;"–"'
DATE = "d mmm yyyy"
SHORT_DATE = "d mmm"

FOREST = "1F4D3A"
INPUT_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
HEAD_FILL = PatternFill("solid", start_color=FOREST, end_color=FOREST)
BAND_FILL = PatternFill("solid", start_color="EEF4F0", end_color="EEF4F0")
TOTAL_FILL = PatternFill("solid", start_color="DDE9E2", end_color="DDE9E2")
THIN = Side(style="thin", color="BFCFC6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass
class TrackerColumn:
    label: str
    start: date
    end: date            # inclusive


def tracker_columns(start: date, period_days: int) -> Tuple[str, List[TrackerColumn]]:
    """
    The spend columns for a period: ("day" | "week", [columns]).

    period_days is the gap to the next payout, exactly what the app stores
    (cycle_end_date = start + period_days). The last column always ends on that
    payout date, which the daily split counts as a day of its own.
    """
    if period_days < 1:
        raise ValueError("period_days must be at least 1")
    payout = start + timedelta(days=period_days)
    if period_days <= 7:
        cols = []
        day = start
        while day <= payout:
            # Not "%-d": that flag does not exist on Windows.
            label = "Payout day" if day == payout else f"{day.strftime('%a')} {day.day} {day.strftime('%b')}"
            cols.append(TrackerColumn(label, day, day))
            day += timedelta(days=1)
        return "day", cols
    weeks = math.ceil(period_days / 7)
    cols = []
    for i in range(weeks):
        col_start = start + timedelta(days=7 * i)
        col_end = payout if i == weeks - 1 else col_start + timedelta(days=6)
        cols.append(TrackerColumn(f"Week {i + 1}", col_start, col_end))
    return "week", cols


def _font(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _input(cell, value, number_format=None):
    cell.value = value
    cell.font = _font(color="0000FF")          # blue = something you type
    cell.fill = INPUT_FILL
    cell.border = BOX
    if number_format:
        cell.number_format = number_format


def _label(cell, text, bold=False, color="000000", italic=False, size=10):
    cell.value = text
    cell.font = _font(bold=bold, color=color, italic=italic, size=size)


def _plain(cell, value, number_format=None, bold=False):
    cell.value = value
    cell.font = _font(bold=bold)
    cell.border = BOX
    if number_format:
        cell.number_format = number_format


def build_budget_workbook(
    categories: Sequence[Tuple[str, Optional[Decimal]]],
    total_amount: Decimal,
    period_days: int,
    start_date: date,
    savings_percentage: Decimal = Decimal("0"),
) -> bytes:
    """Build the workbook and return it as .xlsx bytes."""
    cats = dedupe_categories(categories)
    if not cats:
        raise ValueError("Choose at least one category.")
    kind, columns = tracker_columns(start_date, int(period_days))
    n = len(cats)

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    log = wb.create_sheet("Spend log")

    # ---------------------------------------------------------------- inputs
    _label(ws["A1"], "My budget plan", bold=True, color=FOREST, size=16)
    _label(ws["A2"],
           "Fill in the yellow cells (blue text). Everything else works itself out.",
           italic=True, color="555555")

    rows = {
        "received": 4, "savings_pct": 5, "savings": 6, "spendable": 7,
        "start": 8, "days": 9, "payout": 10, "daycount": 11, "perday": 12,
    }
    _label(ws["A4"], "Money I received (R)", bold=True)
    _input(ws["B4"], float(total_amount), MONEY)
    _label(ws["C4"], "Type what landed in your account for this period.", italic=True, color="777777")

    _label(ws["A5"], "Put aside as savings (%)", bold=True)
    _input(ws["B5"], float(Decimal(savings_percentage) / 100), "0%")
    _label(ws["C5"], "Taken off the top so it is never part of your spending.", italic=True, color="777777")

    _label(ws["A6"], "Savings set aside (R)")
    _plain(ws["B6"], "=ROUND(B4*B5,2)", MONEY)

    _label(ws["A7"], "Available to spend (R)", bold=True)
    _plain(ws["B7"], "=B4-B6", MONEY, bold=True)

    _label(ws["A8"], "Period starts")
    _plain(ws["B8"], start_date, DATE)
    _label(ws["A9"], "Days until next payout")
    _plain(ws["B9"], int(period_days), "0")
    _label(ws["A10"], "Next payout date")
    _plain(ws["B10"], "=B8+B9", DATE)
    _label(ws["A11"], "Days to cover (start day and payout day both count)")
    _plain(ws["B11"], "=B9+1", "0")
    _label(ws["A12"], "Safe to spend per day (R)", bold=True)
    _plain(ws["B12"], "=IF(B11>0,B7/B11,0)", MONEY, bold=True)
    _label(ws["C8"],
           "Fixed when this template was made. Need a different period? Download a new one from the app.",
           italic=True, color="777777")
    _label(ws["C11"], "Same day count the UniWallet daily limit uses.", italic=True, color="777777")

    # ------------------------------------------------------------ table head
    weekly = kind == "week"
    head_row = 16
    from_row, to_row = 14, 15
    first = head_row + 1
    last = head_row + n

    fixed_headers = ["Category", "Planned for the period (R)", "Share of spendable", "Planned per day (R)"]
    if weekly:
        fixed_headers.append("Planned per week (R)")
    col_of_first_track = len(fixed_headers) + 1
    for i, h in enumerate(fixed_headers, start=1):
        c = ws.cell(row=head_row, column=i, value=h)
    for j, col in enumerate(columns):
        c = ws.cell(row=head_row, column=col_of_first_track + j, value=col.label)
    total_col = col_of_first_track + len(columns)
    left_col = total_col + 1
    status_col = total_col + 2
    ws.cell(row=head_row, column=total_col, value="Spent so far (R)")
    ws.cell(row=head_row, column=left_col, value="Left (R)")
    ws.cell(row=head_row, column=status_col, value="Status")
    for c in range(1, status_col + 1):
        cell = ws.cell(row=head_row, column=c)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[head_row].height = 34

    _label(ws.cell(row=from_row, column=col_of_first_track - 1), "From", italic=True, color="777777")
    _label(ws.cell(row=to_row, column=col_of_first_track - 1), "To", italic=True, color="777777")
    ws.cell(row=from_row, column=col_of_first_track - 1).alignment = Alignment(horizontal="right")
    ws.cell(row=to_row, column=col_of_first_track - 1).alignment = Alignment(horizontal="right")
    for j, col in enumerate(columns):
        letter = get_column_letter(col_of_first_track + j)
        f = ws[f"{letter}{from_row}"]
        t = ws[f"{letter}{to_row}"]
        if j == 0:
            f.value = "=$B$8"
        else:
            prev = get_column_letter(col_of_first_track + j - 1)
            f.value = f"={prev}{to_row}+1"
        if kind == "day":
            t.value = f"={letter}{from_row}"
        elif j == len(columns) - 1:
            t.value = "=$B$10"
        else:
            t.value = f"={letter}{from_row}+6"
        for cell in (f, t):
            cell.number_format = SHORT_DATE
            cell.font = _font(color="777777", size=9)
            cell.alignment = Alignment(horizontal="center")

    # ------------------------------------------------------------- category rows
    log_amt = f"'Spend log'!$D${LOG_FIRST_ROW}:$D${LOG_LAST_ROW}"
    log_cat = f"'Spend log'!$C${LOG_FIRST_ROW}:$C${LOG_LAST_ROW}"
    log_date = f"'Spend log'!$A${LOG_FIRST_ROW}:$A${LOG_LAST_ROW}"

    for i, (name, planned) in enumerate(cats):
        r = first + i
        _plain(ws.cell(row=r, column=1), name, bold=True)
        if planned is not None:
            _input(ws.cell(row=r, column=2), float(planned), MONEY)
        else:
            _input(ws.cell(row=r, column=2), None, MONEY)
        _plain(ws.cell(row=r, column=3), f'=IF($B$7>0,B{r}/$B$7,0)', PERCENT)
        _plain(ws.cell(row=r, column=4), f"=B{r}/$B$11", MONEY)
        if weekly:
            _plain(ws.cell(row=r, column=5), f"=D{r}*7", MONEY)
        for j in range(len(columns)):
            letter = get_column_letter(col_of_first_track + j)
            _plain(
                ws.cell(row=r, column=col_of_first_track + j),
                f'=SUMIFS({log_amt},{log_cat},$A{r},{log_date},">="&{letter}${from_row},'
                f'{log_date},"<="&{letter}${to_row})',
                MONEY,
            )
        first_t = get_column_letter(col_of_first_track)
        last_t = get_column_letter(col_of_first_track + len(columns) - 1)
        _plain(ws.cell(row=r, column=total_col), f"=SUM({first_t}{r}:{last_t}{r})", MONEY)
        tl = get_column_letter(total_col)
        _plain(ws.cell(row=r, column=left_col), f'=IF(B{r}="","",B{r}-{tl}{r})', MONEY)
        ll = get_column_letter(left_col)
        _plain(
            ws.cell(row=r, column=status_col),
            f'=IF(B{r}="","Plan me",IF({ll}{r}<0,"Over budget",IF({tl}{r}>=B{r}*0.9,"Nearly used","On track")))',
        )
        ws.cell(row=r, column=status_col).alignment = Alignment(horizontal="center")
        if i % 2 == 1:
            for c in (1, 3, 4) + ((5,) if weekly else ()):
                ws.cell(row=r, column=c).fill = BAND_FILL

    # Other: anything logged under a name that is not in the list above.
    other = last + 1
    total = last + 2
    _plain(ws.cell(row=other, column=1), "Other (not in the list above)", bold=True)
    ws.cell(row=other, column=1).font = _font(italic=True)
    _plain(ws.cell(row=other, column=2), 0, MONEY)
    ws.cell(row=other, column=2).font = _font(color="777777")
    ws.cell(row=other, column=3).border = BOX
    ws.cell(row=other, column=4).border = BOX
    if weekly:
        ws.cell(row=other, column=5).border = BOX
    for j in range(len(columns)):
        letter = get_column_letter(col_of_first_track + j)
        everything = (
            f'SUMIFS({log_amt},{log_date},">="&{letter}${from_row},{log_date},"<="&{letter}${to_row})'
        )
        _plain(
            ws.cell(row=other, column=col_of_first_track + j),
            f"={everything}-SUM({letter}{first}:{letter}{last})",
            MONEY,
        )
    first_t = get_column_letter(col_of_first_track)
    last_t = get_column_letter(col_of_first_track + len(columns) - 1)
    _plain(ws.cell(row=other, column=total_col), f"=SUM({first_t}{other}:{last_t}{other})", MONEY)
    tl = get_column_letter(total_col)
    _plain(ws.cell(row=other, column=left_col), None)
    _plain(ws.cell(row=other, column=status_col), None)

    # Totals
    _plain(ws.cell(row=total, column=1), "Total", bold=True)
    _plain(ws.cell(row=total, column=2), f"=SUM(B{first}:B{last})", MONEY, bold=True)
    _plain(ws.cell(row=total, column=3), f"=IF($B$7>0,B{total}/$B$7,0)", PERCENT, bold=True)
    _plain(ws.cell(row=total, column=4), f"=SUM(D{first}:D{last})", MONEY, bold=True)
    if weekly:
        _plain(ws.cell(row=total, column=5), f"=SUM(E{first}:E{last})", MONEY, bold=True)
    for j in range(len(columns)):
        letter = get_column_letter(col_of_first_track + j)
        _plain(ws.cell(row=total, column=col_of_first_track + j),
               f"=SUM({letter}{first}:{letter}{other})", MONEY, bold=True)
    _plain(ws.cell(row=total, column=total_col), f"=SUM({tl}{first}:{tl}{other})", MONEY, bold=True)
    ll = get_column_letter(left_col)
    _plain(ws.cell(row=total, column=left_col), f"=B{total}-{tl}{total}", MONEY, bold=True)
    _plain(ws.cell(row=total, column=status_col), None)
    for c in range(1, status_col + 1):
        ws.cell(row=total, column=c).fill = TOTAL_FILL

    # Money still to give a job
    gap = total + 2
    _label(ws.cell(row=gap, column=1), "Not planned yet (R)", bold=True)
    _plain(ws.cell(row=gap, column=2), f"=B7-B{total}", MONEY, bold=True)
    _label(ws.cell(row=gap, column=3),
           f'=IF(B{gap}<0,"You have planned more than you can spend — trim a category.",'
           f'IF(B{gap}=0,"Every rand has a job.","Give the rest a job, or leave it as a buffer."))',
           italic=True, color="555555")
    remaining_row = gap + 1
    _label(ws.cell(row=remaining_row, column=1), "Left to spend right now (R)", bold=True)
    _plain(ws.cell(row=remaining_row, column=2), f"=B7-{tl}{total}", MONEY, bold=True)
    _label(ws.cell(row=remaining_row, column=3),
           "Available to spend minus everything in the Spend log.", italic=True, color="777777")

    # Legend
    legend = remaining_row + 2
    notes = [
        "How to use this",
        "1. Yellow cells with blue text are yours to fill in: the money you received, your savings %, and how much you plan for each category.",
        "2. Aim to plan every rand of \"Available to spend\" — the \"Not planned yet\" line shows what is still unassigned.",
        "3. Record each purchase on the Spend log sheet (date, what, category, amount). The columns here add it up for you.",
        f"4. Columns are {'days' if kind == 'day' else 'weeks'} because this budget lasts {int(period_days)} days. The last one ends on payout day, which counts as a day.",
        "5. The same purchase can also go into the UniWallet app, which keeps your daily limit up to date.",
        "Example row on the Spend log sheet shows the format. It is not counted.",
    ]
    for k, text in enumerate(notes):
        _label(ws.cell(row=legend + k, column=1), text, bold=(k == 0), color=FOREST if k == 0 else "333333")

    # Number rows
    for r in (4, 6, 7, 12):
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="right")

    # Conditional formats
    red_font = Font(name=FONT, color="C00000", bold=True)
    red_fill = PatternFill("solid", start_color="F8D7DA", end_color="F8D7DA")
    ws.conditional_formatting.add(
        f"{ll}{first}:{ll}{total}", CellIsRule(operator="lessThan", formula=["0"], font=red_font, fill=red_fill)
    )
    ws.conditional_formatting.add(
        f"B{gap}", CellIsRule(operator="lessThan", formula=["0"], font=red_font, fill=red_fill)
    )
    sl = get_column_letter(status_col)
    ws.conditional_formatting.add(
        f"{sl}{first}:{sl}{last}",
        FormulaRule(formula=[f'{sl}{first}="Over budget"'], font=red_font, fill=red_fill),
    )
    ws.conditional_formatting.add(
        f"{sl}{first}:{sl}{last}",
        FormulaRule(
            formula=[f'{sl}{first}="Nearly used"'],
            font=Font(name=FONT, color="7A5B00", bold=True),
            fill=PatternFill("solid", start_color="FFF3CD", end_color="FFF3CD"),
        ),
    )

    # Widths, freeze
    ws.column_dimensions["A"].width = 44
    for c in range(2, status_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.column_dimensions[get_column_letter(status_col)].width = 14
    ws.freeze_panes = ws.cell(row=head_row + 1, column=2)
    ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------------ log
    _label(log["A1"], "Spend log", bold=True, color=FOREST, size=16)
    _label(log["A2"],
           "One row per purchase. Pick the category from the list (or type your own — it will show under \"Other\").",
           italic=True, color="555555")
    for c, h in enumerate(["Date", "What I bought", "Category", "Amount (R)"], start=1):
        cell = log.cell(row=3, column=c, value=h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BOX
    example = [start_date, "Bread, milk and eggs (example — not counted)", cats[0][0], 85.5]
    for c, v in enumerate(example, start=1):
        cell = log.cell(row=4, column=c, value=v)
        cell.font = _font(italic=True, color="888888")
        cell.border = BOX
    log["A4"].number_format = DATE
    log["D4"].number_format = MONEY
    for r in range(LOG_FIRST_ROW, LOG_FIRST_ROW + 60):
        a, b, c_, d = (log.cell(row=r, column=k) for k in range(1, 5))
        for cell in (a, b, c_, d):
            cell.font = _font(color="0000FF")
            cell.fill = INPUT_FILL
            cell.border = BOX
        a.number_format = DATE
        d.number_format = MONEY
    dv = DataValidation(
        type="list",
        formula1=f"=Budget!$A${first}:$A${last}",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="Not in your list",
        error="That category is not in your plan, so it will count under \"Other\". Continue?",
    )
    log.add_data_validation(dv)
    dv.add(f"C{LOG_FIRST_ROW}:C{LOG_LAST_ROW}")
    dvd = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, errorTitle="Date needed", error="Enter the date of the purchase.")
    log.add_data_validation(dvd)
    dvd.add(f"A{LOG_FIRST_ROW}:A{LOG_LAST_ROW}")
    log.column_dimensions["A"].width = 14
    log.column_dimensions["B"].width = 44
    log.column_dimensions["C"].width = 24
    log.column_dimensions["D"].width = 14
    log.freeze_panes = "A4"
    log.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_filename(start: date, period_days: int) -> str:
    kind = "weekly" if period_days <= 7 else "monthly" if period_days > 21 else "fortnightly"
    return f"uniwallet-{kind}-budget-{start.isoformat()}.xlsx"
